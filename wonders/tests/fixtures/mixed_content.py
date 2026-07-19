"""
Generates small mixed-content fixture files (narrative text + a table + an
embedded/referenced image) for each supported format, at test time — no
binary fixtures checked into the repo. Used by tests/test_ingestion.py to
prove the "no silent content loss" contract from the task spec: a file with
text + table + image must produce chunks representing all three.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import fitz
from docx import Document as DocxDocument

# A valid, minimal 1x1 transparent PNG — no Pillow dependency needed to create it.
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def make_mixed_docx(directory: Path, name: str = "mixed.docx") -> Path:
    document = DocxDocument()
    document.add_paragraph(
        "This is a narrative paragraph about photosynthesis and cellular respiration in plants."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    document.add_picture(io.BytesIO(_TINY_PNG))

    path = directory / name
    document.save(str(path))
    return path


def make_mixed_pdf(directory: Path, name: str = "mixed.pdf") -> Path:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A narrative paragraph about mitochondria and ATP synthesis.")

    x0, y0, cell_w, cell_h = 72, 150, 100, 25
    for r in range(3):
        page.draw_line((x0, y0 + r * cell_h), (x0 + 2 * cell_w, y0 + r * cell_h))
    for c in range(3):
        page.draw_line((x0 + c * cell_w, y0), (x0 + c * cell_w, y0 + 2 * cell_h))
    page.insert_text((x0 + 5, y0 + 17), "Header A")
    page.insert_text((x0 + cell_w + 5, y0 + 17), "Header B")
    page.insert_text((x0 + 5, y0 + 42), "Value 1")
    page.insert_text((x0 + cell_w + 5, y0 + 42), "Value 2")

    page.insert_image(fitz.Rect(72, 250, 122, 300), stream=_TINY_PNG)

    path = directory / name
    pdf.save(str(path))
    pdf.close()
    return path


def make_mixed_markdown(directory: Path, name: str = "mixed.md") -> Path:
    content = (
        "A narrative paragraph about the electron transport chain.\n\n"
        "| Header A | Header B |\n"
        "|----------|----------|\n"
        "| Value 1  | Value 2  |\n\n"
        "![a labeled diagram](diagram.png)\n"
    )
    path = directory / name
    path.write_text(content, encoding="utf-8")
    return path


def make_mixed_txt(directory: Path, name: str = "mixed.txt") -> Path:
    """.txt has no native table/image concept — narrative text only, used to
    prove the plain-text path still works, not the mixed-content contract."""
    content = "A narrative paragraph about ribosomes and protein synthesis.\n\nA second paragraph."
    path = directory / name
    path.write_text(content, encoding="utf-8")
    return path
