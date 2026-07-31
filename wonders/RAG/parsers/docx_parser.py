"""DocxParser — FormatParser implementation for .docx files, via python-docx.

Paragraphs -> NarrativeText, tables -> Table (tab-separated row text),
embedded images -> Image (raw bytes saved to image_output_dir, so the VLM
tool — VLM/imagellm.py's read_image — has a real path to look at) plus, if an
OcrProvider is injected, -> OCRText (the image's literal text — see
RAG/ocr.py).
`iter_inner_content()` (python-docx >= 1.1) preserves document order across
paragraphs and tables, unlike iterating `.paragraphs`/`.tables` separately.
"""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import List, Optional

from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from langchain_core.documents import Document

from RAG.ocr import OcrProvider


class DocxParser:
    def __init__(self, ocr_provider: Optional[OcrProvider] = None):
        self._ocr_provider = ocr_provider

    def parse(self, file_path: str, image_output_dir: Optional[str] = None) -> List[Document]:
        docx_doc = DocxDocument(file_path)
        documents: List[Document] = []

        for item in docx_doc.iter_inner_content():
            if isinstance(item, DocxParagraph):
                text = item.text.strip()
                if text:
                    documents.append(Document(page_content=text, metadata={"category": "NarrativeText"}))
            elif isinstance(item, DocxTable):
                table_text = self._serialize_table(item)
                if table_text:
                    documents.append(Document(page_content=table_text, metadata={"category": "Table"}))

        documents.extend(self._extract_images(docx_doc, file_path, image_output_dir))
        return documents

    @staticmethod
    def _serialize_table(table: DocxTable) -> str:
        rows = ["\t".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        return "\n".join(row for row in rows if row.strip())

    def _extract_images(
        self, docx_doc: DocxDocument, file_path: str, image_output_dir: Optional[str]
    ) -> List[Document]:
        output_dir = Path(image_output_dir or tempfile.gettempdir())
        output_dir.mkdir(parents=True, exist_ok=True)
        stem = Path(file_path).stem
        file_name = Path(file_path).name

        images: List[Document] = []
        index = 0
        for rel in docx_doc.part.rels.values():
            if "image" not in rel.reltype:
                continue
            index += 1
            part = rel.target_part
            ext = Path(part.partname).suffix or ".bin"
            image_path = output_dir / f"{stem}_image_{index}{ext}"
            image_path.write_bytes(part.blob)
            images.append(
                Document(
                    page_content=(
                        f"[Image {index} embedded in {file_name} — use the read_image "
                        f"tool on {image_path} for a visual description.]"
                    ),
                    metadata={"category": "Image", "image_path": str(image_path)},
                )
            )

            ocr_text = self._run_ocr(image_path)
            if ocr_text:
                images.append(
                    Document(
                        page_content=ocr_text,
                        metadata={"category": "OCRText", "image_path": str(image_path)},
                    )
                )
        return images

    def _run_ocr(self, image_path: Path) -> str:
        if self._ocr_provider is None:
            return ""
        try:
            return self._ocr_provider.extract_text(str(image_path)).strip()
        except Exception:  # pragma: no cover - defensive against OCR engine failures
            return ""
